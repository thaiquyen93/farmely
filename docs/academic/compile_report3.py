import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

TARGET_DOCX = "docs/academic/target_srs.docx"
MD_PATH = "docs/academic/SRS_Farmely_Part2_English.md"
OUTPUT_DOCX = "docs/academic/Report3_SRS_Document_Final.docx"

def add_markdown_inlines(paragraph, text, base_font_size=10.0, default_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(base_font_size - 0.5)
            r.font.color.rgb = RGBColor(160, 40, 40)
        else:
            r = paragraph.add_run(token)
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color

def format_cell(cell, text, bold=False, italic=False, font_size=9.0, 
                align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=None, text_color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2.5)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    
    if bold:
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(font_size)
        if text_color:
            r.font.color.rgb = text_color
    else:
        add_markdown_inlines(p, text, base_font_size=font_size, default_color=text_color)
        
    tcPr = cell._tc.get_or_add_tcPr()
    if bg_color:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)

def apply_table_styles(table, header_bg="1E4D78"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="1E4D78"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="1E4D78"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    tblCellMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/>'
        f'<w:right w:w="160" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblCellMar)
    
    trPr0 = table.rows[0]._tr.get_or_add_trPr()
    trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr0.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for r in table.rows[1:]:
        rPr = r._tr.get_or_add_trPr()
        rPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def get_col_widths(headers, total_width=6.4):
    num_cols = len(headers)
    h_text = " ".join(headers).lower()
    
    if num_cols == 7 and "field" in h_text:
        return [0.3, 1.3, 0.7, 0.7, 1.3, 0.7, 1.4]
    elif num_cols == 6 and "field" in h_text:
        return [0.3, 1.5, 0.8, 0.7, 1.5, 1.6]
    elif num_cols == 5 and "field" in h_text:
        return [0.3, 1.6, 1.0, 0.7, 2.8]
    elif num_cols == 3 and "validation" in h_text:
        return [0.8, 1.8, 3.8]
    elif num_cols == 4 and "message" in h_text:
        return [0.8, 2.0, 2.6, 1.0]
    elif num_cols == 5 and ("background" in h_text or "function" in h_text):
        return [0.4, 1.3, 1.7, 1.3, 1.7]
    else:
        w = total_width / num_cols
        return [w] * num_cols

def build_docx_table(doc, headers, data, header_bg="1E4D78"):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    apply_table_styles(table, header_bg=header_bg)
    col_widths = get_col_widths(headers)
    
    # Header row
    for col_idx, h in enumerate(headers):
        format_cell(table.cell(0, col_idx), h, bold=True, font_size=9.0, 
                    align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=header_bg, 
                    text_color=RGBColor(255, 255, 255))
        
    # Data rows
    for row_idx, row in enumerate(data):
        bg = "F8FAFC" if (row_idx % 2 == 1) else "FFFFFF"
        for col_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if (col_idx == 0 or len(str(val)) <= 3 or str(val) in ["X", "✓", "—", "Yes", "No", "None"]) else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(table.cell(row_idx + 1, col_idx), val, bold=False, font_size=8.5, 
                        align=align, bg_color=bg, text_color=RGBColor(40, 40, 40))
            
    if col_widths and len(col_widths) == len(headers):
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def main():
    doc = docx.Document(TARGET_DOCX)
    body = doc.element.body
    
    # Find paragraph 66 ("II. Functional Requirements")
    cut_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if "II. Functional Requirements" in p.text:
            cut_idx = i
            break
            
    print(f"Cutting at paragraph {cut_idx}")
    p_target = doc.paragraphs[cut_idx]._p
    children = list(body)
    target_pos = children.index(p_target)
    print(f"Target body child index: {target_pos} out of {len(children)}")
    
    # Keep the final sectPr at children[-1]
    for el in children[target_pos:-1]:
        body.remove(el)
        
    print(f"Removed placeholder elements. Remaining body elements: {len(body)}")
    
    # Read Markdown
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.splitlines()
    i = 0
    total_lines = len(lines)
    in_code_block = False
    code_lines = []

    while i < total_lines:
        line = lines[i].rstrip()
        stripped = line.strip()

        # Handle code blocks (screen mockups)
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                code_content = "\n".join(code_lines)
                p_box = doc.add_paragraph()
                p_box.paragraph_format.space_before = Pt(4)
                p_box.paragraph_format.space_after = Pt(4)
                p_box.paragraph_format.line_spacing = 1.0
                r = p_box.add_run(code_content)
                r.font.name = "Consolas"
                r.font.size = Pt(8.0)
                r.font.color.rgb = RGBColor(30, 41, 59)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # Major Heading: # II. Functional Requirements
        if stripped.startswith("# II. Functional Requirements"):
            p = doc.add_paragraph(style='Heading 1')
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            r = p.runs[0] if p.runs else p.add_run()
            r.text = "II. Functional Requirements"
            r.font.name = "Calibri"
            r.font.size = Pt(16)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 77, 120)
            i += 1
            continue

        # Feature Heading: ## 1. Feature Name
        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph(style='Heading 2')
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.runs[0] if p.runs else p.add_run()
            r.text = h_text
            r.font.name = "Calibri"
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(30, 77, 120)
            i += 1
            continue

        # Sub-function Heading: ### a. Function Name
        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph(style='Heading 3')
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.runs[0] if p.runs else p.add_run()
            r.text = h_text
            r.font.name = "Calibri"
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            i += 1
            continue

        # Markdown Table Detection
        if stripped.startswith("|") and stripped.endswith("|"):
            raw_table_lines = []
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                raw_table_lines.append(lines[i].strip())
                i += 1

            if len(raw_table_lines) >= 3 and ":---" in raw_table_lines[1]:
                headers = [c.strip() for c in raw_table_lines[0].split("|")[1:-1]]
                data = []
                for row_line in raw_table_lines[2:]:
                    cols = [c.strip() for c in row_line.split("|")[1:-1]]
                    while len(cols) < len(headers):
                        cols.append("")
                    data.append(cols[:len(headers)])

                header_bg = "1E4D78"
                if "field" in " ".join(headers).lower():
                    header_bg = "1E4D78"
                elif "validation" in " ".join(headers).lower():
                    header_bg = "2B6CB0"
                elif "message" in " ".join(headers).lower():
                    header_bg = "3182CE"
                elif "background" in " ".join(headers).lower():
                    header_bg = "1E4D78"

                build_docx_table(doc, headers, data, header_bg=header_bg)
                continue

        # Standard Paragraph / Inlines
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_markdown_inlines(p, stripped, base_font_size=10.0, default_color=RGBColor(30, 41, 59))
        i += 1

    doc.save(OUTPUT_DOCX)
    print(f"Successfully compiled: {OUTPUT_DOCX}")

if __name__ == "__main__":
    main()
