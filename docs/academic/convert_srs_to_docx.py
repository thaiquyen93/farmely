#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Convert SRS_Farmely_v1.md into an official Word document (.docx)
using Template2_SRS Document.docx as base template.
Ensures 100% compliance with FPT ISAD SRS standards.
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

MD_PATH = "D:/isp392/docs/academic/SRS_Farmely_v1.md"
TEMPLATE_PATH = "D:/isp392/docs/academic/Template2_SRS Document.docx"
OUTPUT_PATH = "D:/isp392/docs/academic/SRS_Farmely_Final.docx"
LEGACY_OUTPUT_PATH = "D:/isp392/docs/academic/SRS_Farmely_v1.docx"

CONTEXT_IMG = "D:/isp392/docs/diagrams/00_Context_Diagram_ISAD_Final.png"
OVERVIEW_IMG = "D:/isp392/docs/diagrams/01_System_Overview_v2.png"
ERD_IMG = "D:/isp392/docs/diagrams/erd/00_ERD_Conceptual.png"

def add_markdown_inlines(paragraph, text, base_font_size=10.0, default_color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            r = paragraph.add_run(token[2:-2])
            r.bold = True
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            r = paragraph.add_run(token[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(base_font_size - 0.5)
            r.font.color.rgb = RGBColor(160, 40, 40)
        else:
            r = paragraph.add_run(token)
            r.font.name = "Calibri"
            r.font.size = Pt(base_font_size)
            if default_color:
                r.font.color.rgb = default_color

def format_cell(cell, text, bold=False, italic=False, font_size=9.0, 
                align=WD_ALIGN_PARAGRAPH.LEFT, bg_color=None, text_color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2.5)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    
    if bold:
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Calibri"
        r.font.size = Pt(font_size)
        if text_color:
            r.font.color.rgb = text_color
    else:
        add_markdown_inlines(p, text, base_font_size=font_size, default_color=text_color)
        
    tcPr = cell._tc.get_or_add_tcPr()
    if bg_color:
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_color}"/>')
        tcPr.append(shd)

def apply_table_styles(table, header_bg="1E4D78"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="1E4D78"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="1E4D78"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    tblCellMar = parse_xml(
        f'<w:tblCellMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:left w:w="160" w:type="dxa"/>'
        f'<w:right w:w="160" w:type="dxa"/>'
        f'</w:tblCellMar>'
    )
    tblPr.append(tblCellMar)
    
    trPr0 = table.rows[0]._tr.get_or_add_trPr()
    trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr0.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    for r in table.rows[1:]:
        rPr = r._tr.get_or_add_trPr()
        rPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def get_col_widths(headers, total_width=6.4):
    num_cols = len(headers)
    h_text = " ".join(headers).lower()
    
    if num_cols == 4 and "screen" in h_text and "feature" in h_text:
        return [0.4, 1.3, 1.8, 2.9]
    elif num_cols == 5 and "screen" in h_text and ("cgnn" in h_text or "kh" in h_text):
        return [3.2, 0.8, 0.8, 0.8, 0.8]
    elif num_cols == 3 and ("entity" in h_text or "bảng" in h_text):
        return [0.4, 2.0, 4.0]
    elif num_cols == 3 and "thực thể" in h_text:
        return [1.8, 1.0, 3.6]
    elif num_cols == 4 and "system function" in h_text:
        return [0.4, 1.4, 1.8, 2.8]
    elif num_cols == 7 and "field" in h_text:
        return [0.3, 1.3, 0.7, 0.7, 1.3, 0.7, 1.4]
    elif num_cols == 6 and "field" in h_text:
        return [0.3, 1.5, 0.8, 0.7, 1.5, 1.6]
    elif num_cols == 5 and "field" in h_text:
        return [0.3, 1.6, 1.0, 0.7, 2.8]
    elif num_cols == 3 and "validation" in h_text:
        return [0.8, 1.8, 3.8]
    elif num_cols == 4 and "message" in h_text:
        return [0.8, 2.0, 2.6, 1.0]
    elif num_cols == 4 and "record of changes" in h_text:
        return [1.2, 0.8, 1.2, 3.2]
    else:
        w = total_width / num_cols
        return [w] * num_cols

def build_docx_table(doc, headers, data, header_bg="1E4D78"):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    apply_table_styles(table, header_bg=header_bg)
    col_widths = get_col_widths(headers)
    
    # Header row
    for col_idx, h in enumerate(headers):
        format_cell(table.cell(0, col_idx), h, bold=True, font_size=9.0, 
                    align=WD_ALIGN_PARAGRAPH.CENTER, bg_color=header_bg, 
                    text_color=RGBColor(255, 255, 255))
        
    # Data rows
    for row_idx, row in enumerate(data):
        bg = "F8FAFC" if (row_idx % 2 == 1) else "FFFFFF"
        for col_idx, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if (col_idx == 0 or len(str(val)) <= 3 or str(val) in ["X", "✓", "—", "Có", "Không"]) else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(table.cell(row_idx + 1, col_idx), val, bold=False, font_size=8.5, 
                        align=align, bg_color=bg, text_color=RGBColor(40, 40, 40))
            
    if col_widths and len(col_widths) == len(headers):
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table

def add_image_with_caption(doc, img_path, caption, width_in=6.2):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        r = p_img.add_run()
        r.add_picture(img_path, width=Inches(width_in))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(8)
        rc = p_cap.add_run(caption)
        rc.font.name = "Calibri"
        rc.font.size = Pt(9.5)
        rc.italic = True
        rc.font.color.rgb = RGBColor(90, 90, 90)

def main():
    print(f"Loading template: {TEMPLATE_PATH}")
    doc = docx.Document(TEMPLATE_PATH)
    body = doc.element.body
    
    # Update cover page
    for p in doc.paragraphs[:25]:
        if "<<PROJECT NAME>>" in p.text:
            p.text = "FARMELY – NỀN TẢNG THUÊ PLOT & THEO DÕI TRẢI NGHIỆM CANH TÁC"
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(18)
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = RGBColor(192, 0, 0)
        elif "– Hanoi, August 2022 –" in p.text:
            p.text = "– Hà Nội, Tháng 10/2026 –"
            p.runs[0].font.name = "Calibri"
            p.runs[0].font.size = Pt(11)
            p.runs[0].italic = True

    # Update Table 0 (Record of Changes)
    if len(doc.tables) > 0:
        tbl0 = doc.tables[0]
        if len(tbl0.rows) > 1:
            row1 = tbl0.rows[1]
            row1.cells[0].text = "01/Oct/2026"
            row1.cells[1].text = "A"
            row1.cells[2].text = "NguyênTN"
            row1.cells[3].text = "Khởi tạo SRS Document v1.0 – Hoàn thành toàn bộ Section I & II"
            for c in row1.cells:
                c.paragraphs[0].runs[0].font.size = Pt(9.5)
                c.paragraphs[0].runs[0].font.name = "Calibri"
        for r_idx in range(len(tbl0.rows) - 1, 2, -1):
            tbl0._tbl.remove(tbl0.rows[r_idx]._tr)

    # Remove template placeholder elements from child 30 onwards
    to_remove = [body[i] for i in range(30, len(body) - 1)]
    for el in to_remove:
        body.remove(el)
    print("Template header and TOC preserved. Appending full SRS content...")

    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    lines = md_text.splitlines()
    
    # Skip until "# I. Overview"
    start_idx = 0
    for idx, l in enumerate(lines):
        if l.strip() == "# I. Overview":
            start_idx = idx
            break

    i = start_idx
    total_lines = len(lines)
    in_code_block = False
    code_lines = []
    current_major_section = 1  # 1 = Overview, 2 = Functional Requirements

    while i < total_lines:
        line = lines[i].rstrip()
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                in_code_block = False
                code_content = "\n".join(code_lines)
                if "├── PUBLIC" in code_content:
                    # Site map -> render nicely in a monospace box
                    p_box = doc.add_paragraph()
                    p_box.paragraph_format.space_before = Pt(4)
                    p_box.paragraph_format.space_after = Pt(4)
                    p_box.paragraph_format.line_spacing = 1.05
                    r = p_box.add_run(code_content)
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.0)
                    r.font.color.rgb = RGBColor(40, 40, 40)
                    # Embed Use Case overview diagram right after site map!
                    add_image_with_caption(doc, OVERVIEW_IMG,
                                           "Hình 2.1: Sơ đồ tổng quan phân hệ chức năng hệ thống Farmely (Use Case Diagram v2)",
                                           width_in=6.2)
                else:
                    p_box = doc.add_paragraph()
                    r = p_box.add_run(code_content)
                    r.font.name = "Consolas"
                    r.font.size = Pt(8.5)
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Blank lines
        if not stripped:
            i += 1
            continue

        # Horizontal rules
        if stripped == "---":
            i += 1
            continue

        # Major section 1 or 2
        if stripped.startswith("# "):
            h_text = stripped[2:].strip()
            doc.add_page_break()
            h = doc.add_heading(h_text, level=1)
            h.paragraph_format.keep_with_next = True
            if "II. Functional Requirements" in h_text:
                current_major_section = 2
            else:
                current_major_section = 1
            i += 1
            continue

        if stripped.startswith("## "):
            h_text = stripped[3:].strip()
            h = doc.add_heading(h_text, level=2)
            h.paragraph_format.keep_with_next = True
            i += 1
            continue

        if stripped.startswith("### "):
            h_text = stripped[4:].strip()
            if current_major_section == 1:
                # In Section I.2, a, b, c, d are Heading 4
                if re.match(r'^[a-d]\.\s', h_text):
                    h = doc.add_heading(h_text, level=4)
                else:
                    h = doc.add_heading(h_text, level=3)
            else:
                # In Section II, ALL functions (a, b, c, d, e, f...) are Heading 3!
                h = doc.add_heading(h_text, level=3)
            h.paragraph_format.keep_with_next = True
            i += 1
            continue

        if stripped.startswith("#### "):
            h_text = stripped[5:].strip()
            h = doc.add_heading(h_text, level=4)
            h.paragraph_format.keep_with_next = True
            i += 1
            continue

        # Tables
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl_lines = []
            while i < total_lines and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                tbl_lines.append(lines[i].strip())
                i += 1
            
            if len(tbl_lines) >= 2:
                headers = [c.strip() for c in tbl_lines[0].strip("|").split("|")]
                data_rows = []
                for row_line in tbl_lines[2:]:
                    cells = [c.strip() for c in row_line.strip("|").split("|")]
                    if len(cells) == len(headers):
                        data_rows.append(cells)
                    elif len(cells) < len(headers):
                        cells.extend([""] * (len(headers) - len(cells)))
                        data_rows.append(cells)
                    else:
                        data_rows.append(cells[:len(headers)])
                
                h_joined = " ".join(headers).lower()
                
                # Check if this is the ERD section, embed ERD image right before entities table!
                if "entity" in h_joined or "bảng" in h_joined:
                    add_image_with_caption(doc, ERD_IMG,
                                           "Hình 3.1: Sơ đồ quan hệ thực thể Farmely (Entity Relationship Diagram - 28 thực thể chuẩn 3NF)",
                                           width_in=6.2)
                
                build_docx_table(doc, headers, data_rows)
                
                # If this was the External Entities table in Section 1.2, embed Context Diagram right after it!
                if "thực thể bên ngoài" in h_joined:
                    add_image_with_caption(doc, CONTEXT_IMG,
                                           "Hình 1.1: Sơ đồ ngữ cảnh hệ thống Farmely (Context Diagram - Chuẩn ISAD)",
                                           width_in=6.2)
            continue

        # Blockquote / Callout
        if stripped.startswith("> "):
            bq_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.2)
            add_markdown_inlines(p, bq_text, base_font_size=9.5)
            i += 1
            continue

        # Section II detail sub-headings: *Field Description:*, *Validations:*, *Messages:*, *Additional Logic:*
        if stripped.startswith("*") and stripped.endswith(":*"):
            sub_title = stripped[1:-2]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(sub_title + ":")
            r.bold = True
            r.italic = True
            r.font.name = "Calibri"
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(30, 77, 120)  # Navy blue accent
            i += 1
            continue

        # Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r_bullet = p.add_run("•  ")
            r_bullet.bold = True
            r_bullet.font.name = "Calibri"
            r_bullet.font.size = Pt(10.0)
            r_bullet.font.color.rgb = RGBColor(30, 77, 120)
            add_markdown_inlines(p, bullet_text, base_font_size=10.0)
            i += 1
            continue

        # Numbered list
        match_num = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if match_num:
            num_val = match_num.group(1)
            num_text = match_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            r_num = p.add_run(f"{num_val}. ")
            r_num.bold = True
            r_num.font.name = "Calibri"
            r_num.font.size = Pt(10.0)
            add_markdown_inlines(p, num_text, base_font_size=10.0)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_markdown_inlines(p, stripped, base_font_size=10.0)
        i += 1

    print(f"Saving final document to: {OUTPUT_PATH}")
    doc.save(OUTPUT_PATH)
    print(f"Successfully generated {OUTPUT_PATH}!")
    
    file_size_mb = os.path.getsize(OUTPUT_PATH) / (1024 * 1024)
    print(f"Output file size: {file_size_mb:.2f} MB")

    # Try copying to LEGACY_OUTPUT_PATH if unlocked
    import shutil
    try:
        shutil.copy2(OUTPUT_PATH, LEGACY_OUTPUT_PATH)
        print(f"Also synced to {LEGACY_OUTPUT_PATH} successfully!")
    except Exception as e:
        print(f"Note: {LEGACY_OUTPUT_PATH} is currently open in Word ({e}). Updated file is ready at {OUTPUT_PATH}.")

if __name__ == "__main__":
    main()
